from flask import Flask, request, redirect, url_for, render_template_string, send_file, flash
import sqlite3
from pathlib import Path
from datetime import datetime
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

APP_DIR = Path(__file__).resolve().parent
DB_PATH = APP_DIR / "clienti_salone.db"
app = Flask(__name__)
app.secret_key = "fabio-salone-local"

def connect():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn

def init_db():
    with connect() as con:
        con.execute("""
        CREATE TABLE IF NOT EXISTS clients (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            phone TEXT DEFAULT '',
            birthday TEXT DEFAULT '',
            notes TEXT DEFAULT '',
            created_at TEXT NOT NULL
        )""")
        con.execute("""
        CREATE TABLE IF NOT EXISTS visits (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            client_id INTEGER NOT NULL,
            visit_date TEXT NOT NULL,
            service TEXT NOT NULL,
            price REAL DEFAULT 0,
            brand TEXT DEFAULT '',
            formula TEXT DEFAULT '',
            oxidant TEXT DEFAULT '',
            notes TEXT DEFAULT '',
            FOREIGN KEY(client_id) REFERENCES clients(id) ON DELETE CASCADE
        )""")

init_db()

STYLE = """
:root{--bg:#f5f1ec;--card:#fff;--ink:#2d211b;--gold:#b18a50;--line:#e6ddd3;--muted:#766b63}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--ink);font:16px system-ui,-apple-system,Segoe UI,sans-serif}
header{background:#30231d;color:#fff;padding:20px 16px 16px}header h1{font-size:22px;margin:0 0 5px}header p{margin:0;color:#e8d8c4;font-size:13px}
main{max-width:760px;margin:0 auto;padding:14px}
nav{display:flex;gap:8px;flex-wrap:wrap;margin:12px 0 16px}
a.btn,button,.btn{border:0;border-radius:11px;padding:11px 14px;background:var(--gold);color:white;text-decoration:none;font-weight:650;display:inline-block;font-size:14px;cursor:pointer}
a.secondary{background:#e8dfd6;color:var(--ink)}.card{background:var(--card);padding:16px;border-radius:16px;margin:12px 0;box-shadow:0 2px 10px #2d211b0b;border:1px solid var(--line)}
h2{font-size:18px;margin:0 0 12px}h3{font-size:16px;margin:0 0 6px}
label{display:block;font-size:13px;font-weight:650;margin:10px 0 5px;color:#51433a}
input,textarea,select{width:100%;padding:12px;border:1px solid #d8cec4;border-radius:10px;background:#fff;font:inherit;color:var(--ink)}
textarea{min-height:76px;resize:vertical}.grid{display:grid;grid-template-columns:1fr 1fr;gap:10px}.wide{grid-column:1/-1}
.client{display:flex;justify-content:space-between;align-items:center;gap:12px;padding:12px 0;border-bottom:1px solid var(--line)}.client:last-child{border-bottom:0}
.small{font-size:13px;color:var(--muted)}.price{font-weight:750;color:#74552d}.notice{padding:10px 12px;background:#f1e6d5;border-radius:10px;margin:10px 0;font-size:14px}
table{width:100%;border-collapse:collapse;font-size:14px}td,th{text-align:left;padding:9px 6px;border-bottom:1px solid var(--line);vertical-align:top}
.actions{display:flex;gap:8px;flex-wrap:wrap;margin-top:12px}.empty{color:var(--muted);padding:10px 0}
@media(max-width:460px){.grid{grid-template-columns:1fr}.wide{grid-column:auto}header h1{font-size:20px}main{padding:10px}}
"""

PAGE = """<!doctype html><html lang="it"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>Fabio De Nicola | Clienti</title><style>""" + STYLE + """</style></head><body>
<header><h1>FABIO DE NICOLA</h1><p>Gestionale clienti · colore · servizi · storico</p></header>
<main>
<nav><a class="btn" href="{{ url_for('home') }}">Clienti</a><a class="btn secondary" href="{{ url_for('new_client') }}">+ Nuovo cliente</a><a class="btn secondary" href="{{ url_for('export_excel') }}">Esporta Excel</a></nav>
{% with messages = get_flashed_messages() %}{% if messages %}{% for message in messages %}<div class="notice">{{ message }}</div>{% endfor %}{% endif %}{% endwith %}
{{ body|safe }}
</main></body></html>"""

def page(body, **context):
    return render_template_string(PAGE, body=render_template_string(body, **context))

@app.route("/")
def home():
    q = request.args.get("q", "").strip()
    with connect() as con:
        if q:
            clients = con.execute("""SELECT c.*, COUNT(v.id) visit_count, COALESCE(SUM(v.price),0) total
                FROM clients c LEFT JOIN visits v ON v.client_id=c.id
                WHERE c.name LIKE ? OR c.phone LIKE ?
                GROUP BY c.id ORDER BY c.name""", (f"%{q}%", f"%{q}%")).fetchall()
        else:
            clients = con.execute("""SELECT c.*, COUNT(v.id) visit_count, COALESCE(SUM(v.price),0) total
                FROM clients c LEFT JOIN visits v ON v.client_id=c.id
                GROUP BY c.id ORDER BY c.name""").fetchall()
        total_clients = con.execute("SELECT COUNT(*) FROM clients").fetchone()[0]
        total_revenue = con.execute("SELECT COALESCE(SUM(price),0) FROM visits").fetchone()[0]
    body = """
    <div class="grid">
      <div class="card"><div class="small">Clienti registrati</div><div style="font-size:27px;font-weight:800">{{ total_clients }}</div></div>
      <div class="card"><div class="small">Incasso registrato</div><div style="font-size:27px;font-weight:800">€ {{ '%.2f'|format(total_revenue) }}</div></div>
    </div>
    <div class="card"><h2>Trova un cliente</h2><form method="get"><input name="q" value="{{ q }}" placeholder="Nome o telefono"><div class="actions"><button type="submit">Cerca</button><a class="btn secondary" href="{{ url_for('home') }}">Mostra tutti</a></div></form></div>
    <div class="card"><h2>Elenco clienti</h2>
    {% if clients %}{% for c in clients %}
      <div class="client"><div><strong>{{ c.name }}</strong><div class="small">{{ c.phone or 'Telefono non inserito' }} · {{ c.visit_count }} visite</div><div class="small">Totale speso: € {{ '%.2f'|format(c.total) }}</div></div><a class="btn" href="{{ url_for('client_detail', client_id=c.id) }}">Apri</a></div>
    {% endfor %}{% else %}<div class="empty">Nessun cliente ancora. Premi “+ Nuovo cliente” per iniziare.</div>{% endif %}
    </div>
    <div class="card"><h2>Salva su Excel</h2><p class="small">Scarica un file Excel con l’elenco clienti e lo storico dei servizi.</p><a class="btn" href="{{ url_for('export_excel') }}">Scarica file Excel (.xlsx)</a></div>
    """
    return page(body, clients=clients, total_clients=total_clients, total_revenue=total_revenue, q=q)

@app.route("/client/new", methods=["GET", "POST"])
def new_client():
    if request.method == "POST":
        name = request.form.get("name", "").strip()
        if not name:
            flash("Inserisci almeno il nome del cliente.")
            return redirect(url_for("new_client"))
        with connect() as con:
            con.execute("INSERT INTO clients(name,phone,birthday,notes,created_at) VALUES(?,?,?,?,?)",
                        (name, request.form.get("phone","").strip(), request.form.get("birthday",""),
                         request.form.get("notes","").strip(), datetime.now().isoformat(timespec="minutes")))
        flash("Cliente salvato.")
        return redirect(url_for("home"))
    body = """
    <div class="card"><h2>Nuovo cliente</h2><form method="post">
    <label>Nome e cognome *</label><input name="name" required autocomplete="name">
    <div class="grid"><div><label>Telefono</label><input name="phone" type="tel" autocomplete="tel"></div><div><label>Data di nascita</label><input name="birthday" type="date"></div></div>
    <label>Note sul cliente</label><textarea name="notes" placeholder="Preferenze, allergie segnalate, indicazioni..."></textarea>
    <div class="actions"><button type="submit">Salva cliente</button><a class="btn secondary" href="{{ url_for('home') }}">Annulla</a></div></form></div>
    """
    return page(body)

@app.route("/client/<int:client_id>")
def client_detail(client_id):
    with connect() as con:
        client = con.execute("SELECT * FROM clients WHERE id=?", (client_id,)).fetchone()
        if not client:
            return redirect(url_for("home"))
        visits = con.execute("SELECT * FROM visits WHERE client_id=? ORDER BY visit_date DESC, id DESC", (client_id,)).fetchall()
        total = sum(float(v["price"] or 0) for v in visits)
    body = """
    <div class="card"><h2>{{ client.name }}</h2><div>{{ client.phone or 'Telefono non inserito' }}</div>
    {% if client.birthday %}<div class="small">Nascita: {{ client.birthday }}</div>{% endif %}
    {% if client.notes %}<p>{{ client.notes }}</p>{% endif %}
    <div class="notice">{{ visits|length }} servizi registrati · Totale € {{ '%.2f'|format(total) }}</div>
    {% if client.phone %}<a class="btn secondary" href="https://wa.me/{{ client.phone|replace(' ','')|replace('+','')|replace('-','') }}" target="_blank">Apri WhatsApp</a>{% endif %}
    </div>
    <div class="card"><h2>Registra un servizio</h2><form method="post" action="{{ url_for('add_visit', client_id=client.id) }}">
    <div class="grid"><div><label>Data</label><input type="date" name="visit_date" value="{{ today }}" required></div><div><label>Prezzo (€)</label><input type="number" name="price" min="0" step="0.01" value="0"></div></div>
    <label>Servizio *</label><input name="service" placeholder="Es. colore, taglio, piega" required>
    <div class="grid"><div><label>Marca prodotto</label><input name="brand"></div><div><label>Ossidante</label><input name="oxidant" placeholder="Es. 20 vol"></div></div>
    <label>Formula tecnica / tonalità</label><textarea name="formula" placeholder="Es. radice 6.0 + 6.1..."></textarea>
    <label>Note</label><textarea name="notes"></textarea>
    <div class="actions"><button type="submit">Salva servizio</button></div></form></div>
    <div class="card"><h2>Storico servizi</h2>{% if visits %}<table><tr><th>Data</th><th>Servizio</th><th>Formula</th><th>Prezzo</th></tr>
    {% for v in visits %}<tr><td>{{ v.visit_date }}</td><td>{{ v.service }}<div class="small">{{ v.brand }}{% if v.oxidant %} · {{ v.oxidant }}{% endif %}</div></td><td>{{ v.formula }}<div class="small">{{ v.notes }}</div></td><td>€ {{ '%.2f'|format(v.price or 0) }}</td></tr>{% endfor %}</table>{% else %}<div class="empty">Nessun servizio registrato.</div>{% endif %}</div>
    <div class="actions"><a class="btn secondary" href="{{ url_for('home') }}">← Torna ai clienti</a></div>
    """
    return page(body, client=client, visits=visits, total=total, today=datetime.now().strftime("%Y-%m-%d"))

@app.route("/client/<int:client_id>/visit", methods=["POST"])
def add_visit(client_id):
    service = request.form.get("service", "").strip()
    if not service:
        flash("Inserisci il nome del servizio.")
        return redirect(url_for("client_detail", client_id=client_id))
    try:
        price = float(request.form.get("price") or 0)
    except ValueError:
        price = 0
    with connect() as con:
        exists = con.execute("SELECT id FROM clients WHERE id=?", (client_id,)).fetchone()
        if not exists:
            return redirect(url_for("home"))
        con.execute("""INSERT INTO visits(client_id,visit_date,service,price,brand,formula,oxidant,notes)
        VALUES(?,?,?,?,?,?,?,?)""",
        (client_id, request.form.get("visit_date") or datetime.now().strftime("%Y-%m-%d"), service, price,
         request.form.get("brand","").strip(), request.form.get("formula","").strip(),
         request.form.get("oxidant","").strip(), request.form.get("notes","").strip()))
    flash("Servizio salvato.")
    return redirect(url_for("client_detail", client_id=client_id))

@app.route("/export.xlsx")
def export_excel():
    wb = Workbook()
    ws = wb.active
    ws.title = "Clienti"
    ws.append(["ID", "Nome e cognome", "Telefono", "Data di nascita", "Note", "Data inserimento"])
    with connect() as con:
        clients = con.execute("SELECT * FROM clients ORDER BY name").fetchall()
        visits = con.execute("""SELECT v.id, c.name, c.phone, v.visit_date, v.service, v.price,
            v.brand, v.formula, v.oxidant, v.notes
            FROM visits v JOIN clients c ON c.id=v.client_id
            ORDER BY v.visit_date DESC, c.name""").fetchall()
    for c in clients:
        ws.append([c["id"], c["name"], c["phone"], c["birthday"], c["notes"], c["created_at"]])
    style_sheet(ws)
    wsv = wb.create_sheet("Storico servizi")
    wsv.append(["ID servizio", "Cliente", "Telefono", "Data", "Servizio", "Prezzo (€)", "Marca", "Formula", "Ossidante", "Note"])
    for v in visits:
        wsv.append([v["id"], v["name"], v["phone"], v["visit_date"], v["service"], v["price"], v["brand"], v["formula"], v["oxidant"], v["notes"]])
    style_sheet(wsv)
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    filename = "Fabio_De_Nicola_Clienti_" + datetime.now().strftime("%Y%m%d_%H%M") + ".xlsx"
    return send_file(output, as_attachment=True, download_name=filename,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

def style_sheet(ws):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="30231D")
    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(max_len + 2, 12), 42)

if __name__ == "__main__":
    print("Gestionale avviato. Sul computer apri http://127.0.0.1:5000")
    print("Per il cellulare: usa l'indirizzo IP del computer sulla stessa rete Wi-Fi.")
    app.run(host="0.0.0.0", port=5000, debug=False)
